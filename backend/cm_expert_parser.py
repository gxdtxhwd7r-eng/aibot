"""Safe, header-based parser for CM.Expert XLSX exports."""

from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime
from typing import Any, BinaryIO, Iterator
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", " ", text).strip()


ALIASES: dict[str, set[str]] = {
    "published_at": {"дата публикации", "дата размещения", "опубликовано", "дата объявления"},
    "brand": {"марка", "марка автомобиля"},
    "model": {"модель", "модель автомобиля"},
    "generation": {"поколение", "кузов поколение"},
    "year": {"год", "год выпуска"},
    "mileage": {"пробег", "пробег км"},
    "price": {"цена", "цена руб", "стоимость"},
    "condition": {"состояние", "техническое состояние"},
    "pts": {"птс", "тип птс"},
    "owners": {"количество владельцев", "владельцев", "владельцы"},
    "seller_type": {"тип продавца", "продавец тип"},
    "seller_name": {"имя продавца", "продавец", "контактное лицо"},
    "city": {"город", "регион город", "местоположение"},
    "description": {"комментарий описание продавца", "комментарий продавца", "описание продавца", "описание", "комментарий"},
    "auto_url": {"ссылка auto ru", "auto ru", "ссылка на auto ru", "autor u"},
    "avito_url": {"ссылка avito", "avito", "ссылка на avito"},
    "drom_url": {"ссылка drom", "drom", "ссылка на drom"},
    "cm_url": {"ссылка cm expert", "cm expert", "ссылка в cm expert", "ссылка на cm expert", "ссылка id cm expert", "id cm expert", "cm id"},
    "photos": {"количество фотографий", "фото", "кол во фото"},
}

ALIAS_LOOKUP = {normalize_header(alias): key for key, aliases in ALIASES.items() for alias in aliases}


def header_mapping(headers: list[Any]) -> dict[int, str]:
    return {index: ALIAS_LOOKUP[name] for index, raw in enumerate(headers) if (name := normalize_header(raw)) in ALIAS_LOOKUP}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return result or None


def normalize_url(url: str | None) -> str | None:
    if not url:
        return None
    value = url.strip()
    if not re.match(r"^https?://", value, re.I):
        return value
    parts = urlsplit(value)
    query = urlencode([(k, v) for k, v in parse_qsl(parts.query) if not k.lower().startswith("utm_")])
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip("/"), query, ""))


def parse_row(values: tuple[Any, ...], mapping: dict[int, str]) -> dict[str, Any]:
    row = {field: _text(values[index] if index < len(values) else None) for index, field in mapping.items()}
    source_url = next((normalize_url(row.get(key)) for key in ("avito_url", "auto_url", "drom_url") if row.get(key)), None)
    source = "avito" if row.get("avito_url") else "auto.ru" if row.get("auto_url") else "drom" if row.get("drom_url") else None
    cm_url = normalize_url(row.get("cm_url"))
    identity = cm_url or source_url
    if identity:
        external_id = identity
    else:
        fingerprint = "|".join((row.get(k) or "").lower() for k in ("brand", "model", "year", "mileage", "price", "seller_name", "city", "description"))
        external_id = "hash:" + hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()
    return {**row, "cm_url": cm_url, "source_url": source_url, "source": source, "external_id": external_id}


def iter_xlsx_rows(content: bytes | BinaryIO) -> Iterator[dict[str, Any]]:
    stream = io.BytesIO(content) if isinstance(content, bytes) else content
    workbook = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        mapping = header_mapping(headers)
        if not mapping:
            raise ValueError("Не найдены поддерживаемые заголовки CM.Expert")
        for values in rows:
            yield parse_row(values, mapping)
    finally:
        workbook.close()


def advertisement_text(row: dict[str, Any]) -> str:
    labels = (("Автомобиль", "brand", "model", "generation"), ("Год", "year"), ("Пробег", "mileage"),
              ("Цена", "price"), ("Владельцы", "owners"), ("ПТС", "pts"), ("Состояние", "condition"),
              ("Тип продавца", "seller_type"), ("Город", "city"), ("Описание", "description"))
    lines = []
    for label, *keys in labels:
        value = " ".join(str(row.get(key)) for key in keys if row.get(key))
        if value:
            lines.append(f"{label}: {value}")
    return "\n".join(lines)
